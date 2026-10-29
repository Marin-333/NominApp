"""Script de pruebas exhaustivo — usa BD temporal, no toca nomina.db real."""
import os, sys, tempfile, traceback

TMP = tempfile.mkdtemp(prefix="nomina_test_")
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import paths
paths.APP_DIR = TMP
paths.DB_PATH = os.path.join(TMP, "nomina.db")
paths.REPORTS_DIR = os.path.join(TMP, "reports")
os.makedirs(paths.REPORTS_DIR, exist_ok=True)

import database as db
db.DB_PATH = paths.DB_PATH
db.BACKUP_DIR = os.path.join(TMP, "backups")   # no tocar backups/ reales

import pdf_generator as pdf
pdf.REPORTS_DIR = paths.REPORTS_DIR

import paths as _p
def _grd(empresa_id=None):
    if empresa_id is not None:
        d = os.path.join(paths.REPORTS_DIR, f"empresa_{empresa_id}")
        os.makedirs(d, exist_ok=True)
        return d
    return paths.REPORTS_DIR
_p.get_reports_dir = _grd
db.get_reports_dir = _grd if hasattr(db, 'get_reports_dir') else None
import database
database.get_reports_dir = _grd
pdf.get_reports_dir = _grd

FAIL = []
def check(name, cond, extra=""):
    tag = "OK " if cond else "FAIL"
    print(f"[{tag}] {name} {extra}")
    if not cond: FAIL.append(name)

def run(name, fn):
    try:
        r = fn()
        check(name, True, f"-> {r}" if r is not None else "")
        return r
    except Exception as e:
        check(name, False, f"EXC: {e}")
        traceback.print_exc()
        return None

print("="*70); print("PRUEBAS CAPA DE DATOS"); print("="*70)

run("init_db", db.init_db)
run("migrar_base_datos", db.migrar_base_datos)
run("migrar_base_datos (idempotente)", db.migrar_base_datos)

# ── Cifrado de BD (Fase 2) ─────────────────────────────────────────────────
check("cifrado disponible (SQLCipher)", db.CIFRADO_DB)
if db.CIFRADO_DB:
    check("BD cifrada en disco", db._bd_esta_cifrada(db.DB_PATH))
    check("archivo .key existe", os.path.exists(db.DB_PATH + ".key"))
    check("integrity en BD cifrada", db.verificar_integridad()[0])

# Auth
u = run("autenticar superadmin", lambda: db.autenticar("superadmin", "Admin2024!"))
check("superadmin rol", u and u["rol"] == "superadmin")
check("auth incorrecta", db.autenticar("superadmin", "mala") is None)
db.establecer_sesion(u)  # Sesión requerida por el control de acceso (Fase 3)

# Empresa
ok, msg = db.crear_empresa("Test SA", "900111", "Calle 1", "3001234", periodo_pago="mensual")
check("crear_empresa", ok, msg)
ok2, msg2 = db.crear_empresa("Duplicada", "900111", "", "")
check("NIT duplicado rechazado", not ok2, msg2)
emp = db.listar_empresas()[0]; EID = emp["id"]
run("crear_empresa quincenal", lambda: db.crear_empresa("Quin SA", "900222", "x", "y", periodo_pago="quincenal"))
emp2 = [e for e in db.listar_empresas() if e["nit"]=="900222"][0]; EID2 = emp2["id"]
run("actualizar_empresa", lambda: db.actualizar_empresa(EID2, "Quin SA2", "900222", "d", "t", None, "quincenal"))
check("periodo quincenal guardado", db.obtener_empresa(EID2)["periodo_pago"] == "quincenal")

# Admin
ok, msg = db.crear_admin("admin1", "Pass1234", "Admin Uno", EID)
check("crear_admin", ok, msg)
ok, msg = db.crear_admin("admin_debil", "abc", "Pass Débil", EID)
check("password débil rechazada", not ok)
adm = db.autenticar("admin1", "Pass1234")
check("auth admin", adm is not None and adm["rol"] == "admin")
ok, msg = db.crear_admin("admin1", "DupPass1", "Dup", EID)
check("admin duplicado rechazado", not ok)
run("actualizar_admin", lambda: db.actualizar_admin(adm["id"], "admin1", "Admin Uno B", EID))
run("cambiar_password", lambda: db.cambiar_password(adm["id"], "NewPass2026"))
check("nuevo pass auth", db.autenticar("admin1", "NewPass2026") is not None)
db.establecer_sesion(adm)   # sesión del admin para su propia cuenta
ok, msg = db.cambiar_password_propio(adm["id"], "mala", "NuevaPass1")
check("cambio propio pass incorrecta", not ok)
try:
    db.cambiar_password_propio(u["id"], "x", "OtraPass1"); ajena = False
except PermissionError:
    ajena = True
check("admin NO toca cuenta ajena", ajena)
db.establecer_sesion(u)     # volver a superadmin
ok, msg = db.toggle_activo(adm["id"])
check("toggle admin off", ok and db.autenticar("admin1", "NewPass2026") is None)
db.toggle_activo(adm["id"])

# ── Control de acceso por sesión (Fase 3) ──────────────────────────────────
ok, msg = db.crear_empresa("Otra SA", "900222", "Dir 2", "3009999", periodo_pago="mensual")
EID2 = [e for e in db.listar_empresas() if e["nit"] == "900222"][0]["id"]
db.establecer_sesion(adm)   # admin de EID

def _denegado(fn):
    try:   fn(); return False
    except PermissionError: return True

check("admin NO crea en otra empresa", _denegado(
    lambda: db.crear_empleado(EID2, "X", "999", "Cargo", 1000000, "2024-01-01")))
check("admin NO lista otra empresa", _denegado(
    lambda: db.listar_empleados(EID2)))
check("admin NO crea empresas", _denegado(
    lambda: db.crear_empresa("X SA", "1", "d", "t")))
check("admin NO crea admins", _denegado(
    lambda: db.crear_admin("x2", "Pass1234", "X", EID)))
check("admin NO resetea passwords", _denegado(
    lambda: db.cambiar_password(adm["id"], "Pass1234")))
check("admin NO exporta backup completo", _denegado(
    lambda: db.exportar_backup(os.devnull)))
check("admin NO ve bitácora global", _denegado(
    lambda: db.listar_bitacora(50)))
check("admin SÍ opera su empresa", isinstance(db.listar_empleados(EID), list))
check("admin SÍ ve su bitácora", isinstance(db.listar_bitacora(10, empresa_id=EID), list))
check("sin sesión = sin acceso", (db.cerrar_sesion() or True) and _denegado(
    lambda: db.listar_empleados(EID)))
db.establecer_sesion(u)   # volver a superadmin para el resto del test

# ── Exactitud decimal en montos (Fase 2) ───────────────────────────────────
calc_dec = db.calcular_nomina(1234567.89, 13, 2.5, 15432.10, 333333.33, 11111.11,
                              periodo_pago='mensual', prestamo=5000.55)
check("decimal: bruto-neto exacto",
      round(calc_dec["salario_bruto"] - calc_dec["salud"] - calc_dec["pension"]
            - 11111.11 - 5000.55, 2) == calc_dec["salario_neto"])
check("decimal: resultados a 2 decimales", all(
    round(v, 2) == v for v in calc_dec.values() if isinstance(v, float)))
calc_prest = db.calcular_pago_prestacion(123456.78, 7, 100.5, 50.25, 0)
check("decimal prestacion",
      calc_prest["salario_neto"] ==
      round(calc_prest["salario_bruto"] - 50.25, 2))

# Empleados
ok, msg = db.crear_empleado(EID, "Juan Perez", "1001", "Operario", 2000000, "2024-01-15",
                            "Ahorros", "123456", 35, "salario", None, "juan@test.com")
check("crear_empleado salario", ok, msg)
ok, msg = db.crear_empleado(EID, "Juan Dup", "1001", "X", 1, "")
check("doc duplicado rechazado", not ok)
ok, msg = db.crear_empleado(EID, "Ana Prest", "2002", "Consultora", 150000, "2024-03-01",
                            None, None, 30, "prestacion", 150000, None)
check("crear_empleado prestacion", ok, msg)
emps = db.listar_empleados(EID)
E1 = next(e for e in emps if e["documento"] == "1001")["id"]
E2 = next(e for e in emps if e["documento"] == "2002")["id"]
check("tipo_vinculacion", db.obtener_empleado(E2, EID)["tipo_vinculacion"] == "prestacion")
run("actualizar_empleado", lambda: db.actualizar_empleado(E1, EID, "Juan P", "1001", "Op", 2100000, "2024-01-15"))
check("salario actualizado", db.obtener_empleado(E1, EID)["salario"] == 2100000)

# Cálculo nómina
c = db.calcular_nomina(2000000, 30, 0, 0, 0, 0, periodo_pago='mensual')
print("   mensual:", c)
check("mensual bruto == salario", c["salario_bruto"] == 2000000)
c1 = db.calcular_nomina(2000000, 15, 0, 0, 0, 0, periodo_pago='quincenal', quincena=1)
c2 = db.calcular_nomina(2000000, 15, 0, 0, 0, 0, periodo_pago='quincenal', quincena=2)
print("   Q1:", c1, "\n   Q2:", c2)
check("Q1 salud>0 pension=0", c1["salud"] > 0 and c1["pension"] == 0)
check("Q2 pension>0 salud=0", c2["pension"] > 0 and c2["salud"] == 0)
check("Q1+Q2 ded == mensual ded",
      abs((c1["salud"]+c1["pension"] + c2["salud"]+c2["pension"]) - (c["salud"]+c["pension"])) < 0.01,
      f"Q={c1['salud']+c1['pension']+c2['salud']+c2['pension']} M={c['salud']+c['pension']}")
check("salud y pension mensual etiquetadas",
      c["salud"] == db.DEDUCCION_SALUD_QUINCENA and c["pension"] == db.DEDUCCION_PENSION_QUINCENA,
      f"salud={c['salud']} pension={c['pension']}")

# Guardar nómina
ok, msg, nid = db.guardar_nomina(EID, E1, "2026-09", 30, 4, 20000, 100000, 50000, "obs", 2100000)
check("guardar_nomina mensual", ok, msg)
ok, msg, nid2 = db.guardar_nomina(EID, E1, "2026-09", 30, 0, 0, 0, 0, "", 2100000)
check("nómina duplicada rechazada", not ok and nid2 == nid, msg)
# Empleado quincenal en empresa 2
ok, msg = db.crear_empleado(EID2, "Luis Quin", "3003", "Vendedor", 1800000, "2024-02-01")
E3 = db.listar_empleados(EID2)[0]["id"]
ok, msg, nidq = db.guardar_nomina(EID2, E3, "2026-09", 15, 0,0,0,0,"",1800000, quincena=1, periodo_pago='quincenal')
check("guardar_nomina quincenal Q1", ok, msg)
nq1 = db.obtener_nomina(nidq, EID2)
check("Q1 salud descontada", nq1["salud"] == db.DEDUCCION_SALUD_QUINCENA, f"salud={nq1['salud']}")
n = db.obtener_nomina(nid, EID)
check("nomina datos", n["empleado_nombre"] == "Juan P" and n["periodo"] == "2026-09")

# Prestación
ok, msg, nidp = db.guardar_nomina_prestacion(EID, E2, "2026-09", 150000, 12, 50000, 10000, 0)
check("guardar_nomina_prestacion", ok, msg)
np_ = db.obtener_nomina(nidp, EID)
check("prestacion bruto", np_["salario_bruto"] == 150000*12 + 50000, f"{np_['salario_bruto']}")
check("prestacion sin deducciones ley", np_["salud"] == 0 and np_["pension"] == 0)

# Contratos
ok, msg = db.crear_contrato(EID, E2, "Contrato obra", 5000000, "2026-09-01", "2026-12-31", 100000, "obs")
check("crear_contrato", ok, msg)
cont = db.listar_contratos(EID)[0]; CID = cont["id"]
ok, msg, nidc = db.liquidar_contrato(CID, EID, descontar_prestamos=False)
check("liquidar_contrato", ok, msg)
check("contrato liquidado", db.obtener_contrato(CID, EID)["estado"] == "liquidado")
ok, msg, _ = db.liquidar_contrato(CID, EID)
check("re-liquidación rechazada", not ok)
ok, msg = db.eliminar_contrato(CID, EID)
check("no eliminar contrato liquidado", not ok, msg)

# Préstamos
ok, msg = db.crear_prestamo(EID, E1, 500000, "2026-09-01", "prestamo test")
check("crear_prestamo", ok, msg)
pres = db.listar_prestamos(EID)[0]
check("saldo inicial", pres["saldo"] == 500000)
ok, msg = db.abonar_prestamo(pres["id"], 200000)
check("abonar_prestamo", ok, msg)
check("saldo tras abono", db.listar_prestamos(EID)[0]["saldo"] == 300000)
ok, msg = db.abonar_prestamo(pres["id"], 999999)
check("abono mayor que saldo rechazado", not ok)
ok, msg = db.abonar_prestamo(pres["id"], 300000)
check("prestamo pagado", db.listar_prestamos(EID)[0]["estado"] == "pagado")

# nómina con préstamo activo
db.crear_prestamo(EID, E1, 100000, "2026-09-05", "p2")
ok, msg, nid3 = db.guardar_nomina(EID, E1, "2026-10", 30, 0,0,0,0,"",2100000, prestamo=40000)
check("nomina con prestamo", ok, msg)
p2 = [p for p in db.listar_prestamos(EID) if p["monto"] == 100000][0]
check("saldo prestamo descontado", p2["saldo"] == 60000, f"saldo={p2['saldo']}")

# Bitácora
db.registrar_bitacora(adm["id"], "Admin Uno B", "Prueba", "detalle test", empresa_id=EID)
check("bitacora", len(db.listar_bitacora(10, empresa_id=EID)) >= 1)

# Excel
xlsx = os.path.join(TMP, "test.xlsx")
run("exportar_nomina_excel", lambda: db.exportar_nomina_excel(EID, "2026-09", xlsx))
check("excel existe", os.path.exists(xlsx))
import openpyxl
wb = openpyxl.load_workbook(xlsx); ws = wb.active
check("excel datos", ws.cell(row=5, column=2).value is not None, f"A5={ws.cell(row=5,column=2).value}")

# Backup empresa + restore
bk = os.path.join(TMP, "backup_emp.db")
ok, msg = db.exportar_backup_empresa(EID, bk)
check("exportar_backup_empresa", ok, msg)
check("backup empresa cifrado", db._bd_esta_cifrada(bk) if db.CIFRADO_DB else True)
bc = db._conectar_archivo(bk)   # el export queda cifrado con la llave interna
tables = [r[0] for r in bc.execute("SELECT name FROM sqlite_master WHERE type='table'")]
print("   tablas en backup:", tables)
check("contratos en backup", "contratos" in tables)
ecols = [r[1] for r in bc.execute("PRAGMA table_info(empleados)")]
check("tarifa_diaria en backup", "tarifa_diaria" in ecols, str(ecols))
check("correo en backup", "correo" in ecols)
nc_bk = bc.execute("SELECT COUNT(*) FROM contratos").fetchone()[0] if "contratos" in tables else -1
print(f"   contratos en backup: {nc_bk}")
bc.close()

ok, msg = db.importar_backup_empresa(EID, bk)
check("importar_backup_empresa", ok, msg)
check("contratos restaurados", len(db.listar_contratos(EID)) == 1)

# Backup completo
bk2 = os.path.join(TMP, "full.db")
run("exportar_backup", lambda: db.exportar_backup(bk2))
check("backup completo cifrado", db._bd_esta_cifrada(bk2) if db.CIFRADO_DB else True)

# ── Backup con contraseña (Fase 2, portabilidad) ───────────────────────────
if db.CIFRADO_DB:
    bkp = os.path.join(TMP, "backup_pw.db")
    ok, msg = db.exportar_backup(bkp, password="BkpPass123")
    check("export con password", ok, msg)
    check("backup pw cifrado", db._bd_esta_cifrada(bkp))
    check("abre con password", db.probar_clave_db(bkp, "BkpPass123"))
    check("rechaza password malo", not db.probar_clave_db(bkp, "mala"))
    check("rechaza llave interna", not db.probar_clave_db(bkp))
    bke = os.path.join(TMP, "backup_emp_pw.db")
    ok, msg = db.exportar_backup_empresa(EID, bke, password="EmpPass1")
    check("export empresa con password", ok, msg)
    ok, msg = db.importar_backup_empresa(EID, bke, password="EmpPass1")
    check("import empresa con password", ok, msg)

print("="*70); print("PRUEBAS PDF"); print("="*70)

# Tras el restore los ids cambiaron — re-obtener nóminas por quincena
_noms_all = db.listar_nominas(EID)
nid  = next((n["id"] for n in _noms_all if n.get("quincena")==0),  None)
nidp = next((n["id"] for n in _noms_all if n.get("quincena")==98), None)
nidc = next((n["id"] for n in _noms_all if n.get("quincena")==99), None)

nom_full = db.obtener_nomina(nid, EID) if nid else None
if nom_full:
    run("generar_colilla mensual", lambda: pdf.generar_colilla(nom_full, empresa_id=EID))
else:
    check("generar_colilla mensual", False, "nómina no encontrada tras restore")

nom_p = db.obtener_nomina(nidp, EID) if nidp else None
if nom_p:
    run("generar_colilla prestacion", lambda: pdf.generar_colilla(nom_p, empresa_id=EID))
    check("colilla prestacion empleado", nom_p["empleado_nombre"] is not None)
else:
    check("generar_colilla prestacion", False, "nómina prestación no encontrada tras restore")

nom_c = db.obtener_nomina(nidc, EID) if nidc else None
if nom_c:
    run("generar_colilla contrato", lambda: pdf.generar_colilla(nom_c, empresa_id=EID))
else:
    check("generar_colilla contrato", False, "nómina contrato no encontrada tras restore")

noms = db.listar_nominas(EID, "2026-09")
run("generar_reporte_nomina", lambda: pdf.generar_reporte_nomina(emp, noms, "2026-09", empresa_id=EID))
emps_rep = db.obtener_reporte_empleados(EID, True)
run("generar_reporte_empleados", lambda: pdf.generar_reporte_empleados(emps_rep, emp, empresa_id=EID))
regs = db.listar_bitacora(200, empresa_id=EID)
run("generar_reporte_bitacora", lambda: pdf.generar_reporte_bitacora(regs, emp, empresa_id=EID))
run("generar_reporte_bitacora superadmin", lambda: pdf.generar_reporte_bitacora(regs, empresa_id=None))
emp_e = db.obtener_empleado(E1, EID) or {"nombre":"X","documento":"0","cargo":"c","salario":1}
run("generar_certificado_laboral", lambda: pdf.generar_certificado_laboral(emp_e, emp, empresa_id=EID))

print("="*70); print("ELIMINACIÓN DEFINITIVA"); print("="*70)
# ids frescos tras el restore (remap)
E1 = next(e["id"] for e in db.listar_empleados(EID, True) if e["documento"] == "1001")
E2 = next(e["id"] for e in db.listar_empleados(EID, True) if e["documento"] == "2002")
ok, msg = db.eliminar_empleado_definitivo(E1, EID, eliminar_nominas=False)
check("rechaza eliminar con nóminas si no se pide", not ok, msg)
ok, msg = db.eliminar_empleado_definitivo(E2, EID, eliminar_nominas=True)
check("eliminar definitivo", ok, msg)
check("empleado borrado", db.obtener_empleado(E2, EID) is None)

pdfs = [f for f in os.listdir(paths.REPORTS_DIR) if f.endswith(".pdf")]
pdfs += [os.path.join(r,f) for r,d,fs in os.walk(paths.REPORTS_DIR) for f in fs if f.endswith(".pdf")]
print("   PDFs generados:", len(set(pdfs)))

print("="*70)
print(f"RESULTADO: {len(FAIL)} fallos" + (f"  ->  {FAIL}" if FAIL else "  (todo OK)"))
sys.exit(1 if FAIL else 0)
